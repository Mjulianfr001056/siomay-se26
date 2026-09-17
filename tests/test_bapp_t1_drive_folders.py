"""Regression tests for BAPP Termin 1 Drive collection insertion."""

import io
import unittest
from unittest.mock import patch

from docx import Document
from PIL import Image

from src import bapp_pml, bapp_ppl


MODULES = (bapp_ppl, bapp_pml)


def _downloaded_image(_file_id):
    image = Image.new("RGB", (1200, 800), color=(40, 100, 180))
    stream = io.BytesIO()
    image.save(stream, format="PNG")
    stream.seek(0)
    return stream, image


class BappTermin1DriveFolderTests(unittest.TestCase):
    def test_more_than_five_images_are_paginated_instead_of_truncated(self):
        links = ",".join(
            f"https://drive.google.com/file/d/image-{number}/view"
            for number in range(1, 7)
        )
        for module in MODULES:
            with self.subTest(module=module.__name__):
                document = Document()
                paragraph = document.add_paragraph()
                paragraph.add_run(module.BUKTI_PLACEHOLDER[:12])
                paragraph.add_run(module.BUKTI_PLACEHOLDER[12:])

                with patch.object(
                    module, "_download_drive_image", side_effect=_downloaded_image
                ) as downloader:
                    count, warnings = module.insert_gdrive_images(document, links)

                self.assertEqual(count, 6)
                self.assertEqual(warnings, [])
                self.assertEqual(downloader.call_count, 6)
                self.assertEqual(len(document.inline_shapes), 6)
                self.assertEqual(
                    len(document.element.body.xpath('.//w:br[@w:type="page"]')),
                    1,
                )
                self.assertNotIn(module.BUKTI_PLACEHOLDER, paragraph.text)


if __name__ == "__main__":
    unittest.main()