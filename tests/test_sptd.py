"""Tests for SPTD validation, derived values, and DOCX formatting."""

import os
import tempfile
import unittest

from docx import Document
from docx.dml.color import RGBColor
import pandas as pd

from src import sptd
from src.document_generator import _iter_paragraphs
from src.workflow import get_document_by_id


class SptdTests(unittest.TestCase):
    def _write_input(self, path, rows, columns=None):
        dataframe = pd.DataFrame(rows, columns=columns)
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            dataframe.to_excel(writer, sheet_name=sptd.SHEET_NAME, index=False)

    def _row(self, **overrides):
        row = {
            "nama_ppl": "Niki Zefanya",
            "nama_pml": "Arash Buana",
            "kecamatan": "TABUNGANEN",
            "jml_keluarga": "17",
            "jml_usaha": "9",
            "persentase": "2.32",
        }
        row.update(overrides)
        return row

    def test_bundled_input_validates_and_has_three_populated_rows(self):
        document = get_document_by_id("sptd")
        ok, errors, dfs = sptd.validate_input(document.input_template_path)

        self.assertTrue(ok, errors)
        self.assertEqual(len(dfs[sptd.SHEET_NAME]), 3)

    def test_blank_counts_are_zero_but_non_numeric_counts_are_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            blank_path = os.path.join(directory, "blank.xlsx")
            self._write_input(blank_path, [self._row(
                jml_keluarga="", jml_usaha="",
            )])
            ok, errors, _ = sptd.validate_input(blank_path)
            self.assertTrue(ok, errors)

            invalid_path = os.path.join(directory, "invalid.xlsx")
            self._write_input(invalid_path, [self._row(jml_usaha="sembilan")])
            ok, errors, _ = sptd.validate_input(invalid_path)
            self.assertFalse(ok)
            self.assertIn("jml_usaha", errors[0])
            self.assertIn("Baris 2", errors[0])

    def test_missing_required_column_is_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            path = os.path.join(directory, "missing.xlsx")
            row = self._row()
            del row["persentase"]
            self._write_input(path, [row])

            ok, errors, _ = sptd.validate_input(path)

        self.assertFalse(ok)
        self.assertIn("persentase", errors[0])

    def test_generation_calculates_responden_and_preserves_template_format(self):
        document_type = get_document_by_id("sptd")
        ok, errors, dfs = sptd.validate_input(document_type.input_template_path)
        self.assertTrue(ok, errors)

        with tempfile.TemporaryDirectory() as directory:
            events = list(sptd.iter_generate(
                dfs, document_type.builtin_template_path, directory,
            ))
            paths = [event["path"] for event in events if event["t"] == "file"]
            self.assertEqual(len(paths), 3)
            for path, (_, row) in zip(
                    paths, dfs[sptd.SHEET_NAME].iterrows()):
                generated = Document(path)
                generated_text = "\n".join(
                    paragraph.text for paragraph in _iter_paragraphs(generated)
                )
                total = sptd._format_decimal(
                    sptd._count_value(
                        row["jml_keluarga"], row_number=0,
                        column="jml_keluarga",
                    )
                    + sptd._count_value(
                        row["jml_usaha"], row_number=0,
                        column="jml_usaha",
                    )
                )
                self.assertIn(str(row["nama_ppl"]), generated_text)
                self.assertIn(str(row["nama_pml"]), generated_text)
                self.assertIn(str(row["kecamatan"]), generated_text)
                self.assertIn(total, generated_text)
                self.assertNotIn("{{", generated_text)

            output = Document(paths[0])

            paragraphs = list(_iter_paragraphs(output))
            all_text = "\n".join(paragraph.text for paragraph in paragraphs)
            self.assertIn("sebanyak 17  keluarga dan 9 usaha", all_text)
            self.assertIn("sebanyak 26 responden", all_text)
            self.assertIn("2.32%", all_text)
            for placeholder in sptd.BUILTIN_FIELDS:
                self.assertNotIn("{{" + placeholder + "}}", all_text)

            expected = {
                "17": True,
                "9": True,
                "26": True,
                "2.32": True,
                "Niki Zefanya": False,
                "Arash Buana": False,
                "TABUNGANEN": False,
            }
            for value, bold in expected.items():
                matching_runs = [
                    run for paragraph in paragraphs for run in paragraph.runs
                    if run.text == value
                ]
                self.assertTrue(matching_runs, value)
                for run in matching_runs:
                    self.assertEqual(bool(run.bold), bold, value)
                    self.assertEqual(
                        run.font.color.rgb, RGBColor(0x00, 0x00, 0x00), value
                    )

    def test_generation_treats_blank_counts_as_zero(self):
        with tempfile.TemporaryDirectory() as directory:
            template_path = os.path.join(directory, "template.docx")
            template = Document()
            paragraph = template.add_paragraph()
            for text in (
                "Keluarga ", "{{jml_keluarga}}", ", usaha ",
                "{{jml_usaha}}", ", total ", "{{responden}}", ".",
            ):
                run = paragraph.add_run(text)
                if text.startswith("{{"):
                    run.bold = True
            template.save(template_path)
            dfs = {sptd.SHEET_NAME: pd.DataFrame([
                self._row(jml_keluarga="", jml_usaha="")
            ])}

            events = list(sptd.iter_generate(dfs, template_path, directory))
            output_path = next(
                event["path"] for event in events if event["t"] == "file"
            )
            output = Document(output_path)

        self.assertEqual(
            output.paragraphs[0].text,
            "Keluarga 0, usaha 0, total 0.",
        )
        zero_runs = [run for run in output.paragraphs[0].runs if run.text == "0"]
        self.assertEqual(len(zero_runs), 3)
        for zero_run in zero_runs:
            self.assertTrue(zero_run.bold)
            self.assertEqual(zero_run.font.color.rgb, RGBColor(0, 0, 0))

    def test_replacement_does_not_recolor_surrounding_text_in_same_run(self):
        document = Document()
        run = document.add_paragraph().add_run("Merah {{responden}} tetap merah")
        run.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)
        run.bold = True

        sptd.replace_text_preserving_runs(
            document, {"{{responden}}": "26"}
        )

        runs = document.paragraphs[0].runs
        self.assertEqual(document.paragraphs[0].text, "Merah 26 tetap merah")
        self.assertEqual(runs[0].font.color.rgb, RGBColor(0xFF, 0x00, 0x00))
        self.assertEqual(runs[1].text, "26")
        self.assertTrue(runs[1].bold)
        self.assertEqual(runs[1].font.color.rgb, RGBColor(0x00, 0x00, 0x00))
        self.assertEqual(runs[2].font.color.rgb, RGBColor(0xFF, 0x00, 0x00))


if __name__ == "__main__":
    unittest.main()