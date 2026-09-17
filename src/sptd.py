"""Generator Surat Pernyataan Tidak Ditemukan (SPTD) SE2026."""

from decimal import Decimal, InvalidOperation
from copy import deepcopy
import os
import re

import pandas as pd
from docx import Document
from docx.dml.color import RGBColor
from docx.text.run import Run

from src.document_generator import (
    clean_value,
    row_placeholder_replacements,
    validate_custom_columns,
)


SHEET_NAME = "data_mitra"

REQUIRED_COLUMNS = [
    "nama_ppl",
    "nama_pml",
    "kecamatan",
    "jml_keluarga",
    "jml_usaha",
    "persentase",
]

BUILTIN_FIELDS = set(REQUIRED_COLUMNS) | {"responden"}


def _populated_rows(dataframe: pd.DataFrame) -> pd.DataFrame:
    """Remove rows that contain no input value, including formatted XLSX rows."""
    if dataframe.empty:
        return dataframe
    populated = dataframe.apply(
        lambda row: any(clean_value(value) for value in row), axis=1
    )
    return dataframe.loc[populated].copy()


def _count_value(value, *, row_number: int, column: str) -> Decimal:
    """Return a respondent count, treating blanks as zero."""
    text = clean_value(value)
    if not text:
        return Decimal(0)
    try:
        number = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(
            f"Baris {row_number}, kolom '{column}' harus berupa angka."
        ) from exc
    if not number.is_finite():
        raise ValueError(
            f"Baris {row_number}, kolom '{column}' harus berupa angka."
        )
    return number


def _format_decimal(value: Decimal) -> str:
    """Format a decimal without scientific notation or redundant zeroes."""
    result = format(value, "f")
    if "." in result:
        result = result.rstrip("0").rstrip(".")
    return "0" if result in {"", "-0"} else result


def validate_input(file_path: str, custom_fields=None):
    """Validate an SPTD workbook and return ``(is_valid, errors, dfs)``."""
    errors, dfs = [], {}
    try:
        workbook = pd.ExcelFile(file_path)
    except Exception as exc:
        return False, [f"Gagal membaca file Excel: {exc}"], {}

    try:
        if SHEET_NAME not in workbook.sheet_names:
            return False, [
                f"Sheet '{SHEET_NAME}' tidak ditemukan. Sheet yang tersedia: "
                + ", ".join(workbook.sheet_names)
            ], {}
        dataframe = pd.read_excel(workbook, sheet_name=SHEET_NAME, dtype=str)
    finally:
        workbook.close()

    dataframe.columns = [str(column).strip() for column in dataframe.columns]
    dataframe = _populated_rows(dataframe.fillna(""))
    dfs[SHEET_NAME] = dataframe

    missing = [column for column in REQUIRED_COLUMNS if column not in dataframe.columns]
    if missing:
        errors.append(
            f"Sheet '{SHEET_NAME}' kehilangan kolom: {', '.join(missing)}"
        )
        return False, errors, dfs

    if dataframe.empty:
        errors.append(f"Sheet '{SHEET_NAME}' kosong.")
        return False, errors, dfs

    for index, row in dataframe.iterrows():
        excel_row = index + 2
        for column in ("jml_keluarga", "jml_usaha"):
            try:
                _count_value(
                    row.get(column, ""), row_number=excel_row, column=column
                )
            except ValueError as exc:
                errors.append(str(exc))

    errors.extend(validate_custom_columns(dfs, SHEET_NAME, custom_fields))
    return not errors, errors, dfs


def replace_text_preserving_runs(doc: Document, replacements: dict) -> None:
    """Replace split-run placeholders, preserving bold and forcing black text."""
    token_pattern = re.compile(r"\{\{[^}]+\}\}")

    def process_paragraphs(paragraphs):
        for paragraph in paragraphs:
            while paragraph.runs:
                runs = paragraph.runs
                full_text = "".join(run.text for run in runs)
                match = next(
                    (match for match in token_pattern.finditer(full_text)
                     if match.group(0) in replacements),
                    None,
                )
                if match is None:
                    break

                starts, position = [], 0
                for run in runs:
                    starts.append(position)
                    position += len(run.text)
                first = max(i for i, start in enumerate(starts)
                            if start <= match.start())
                last = max(i for i, start in enumerate(starts)
                           if start < match.end())
                prefix = runs[first].text[:match.start() - starts[first]]
                suffix = runs[last].text[match.end() - starts[last]:]

                # Isolate the inserted value in a clone of the placeholder's
                # first run. This preserves its bold state while allowing its
                # colour to be set without recolouring surrounding text.
                replacement_xml = deepcopy(runs[first]._r)
                replacement_run = Run(replacement_xml, paragraph)
                replacement_run.text = str(replacements[match.group(0)])
                replacement_run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

                if first == last:
                    suffix_xml = deepcopy(runs[first]._r)
                    suffix_run = Run(suffix_xml, paragraph)
                    suffix_run.text = suffix
                    runs[first].text = prefix
                    runs[first]._r.addnext(suffix_xml)
                    runs[first]._r.addnext(replacement_xml)
                else:
                    runs[first].text = prefix
                    runs[last].text = suffix
                    for run_index in range(first + 1, last):
                        runs[run_index].text = ""
                    runs[first]._r.addnext(replacement_xml)

    def process_table(table):
        for row in table.rows:
            for cell in row.cells:
                process_paragraphs(cell.paragraphs)
                for nested_table in cell.tables:
                    process_table(nested_table)

    process_paragraphs(doc.paragraphs)
    for table in doc.tables:
        process_table(table)
    seen = set()
    for section in doc.sections:
        for story in (
            section.header, section.footer,
            section.first_page_header, section.first_page_footer,
            section.even_page_header, section.even_page_footer,
        ):
            story_id = id(story._element)
            if story_id in seen:
                continue
            seen.add(story_id)
            process_paragraphs(story.paragraphs)
            for table in story.tables:
                process_table(table)


def _slug(value: str) -> str:
    result = re.sub(r"[^A-Za-z0-9]+", "_", value).strip("_")
    return result[:40] or "tanpa_nama"


def iter_generate(dfs: dict, template_path: str, out_dir: str):
    """Generate one populated SPTD document for each non-empty input row."""
    os.makedirs(out_dir, exist_ok=True)
    dataframe = dfs.get(SHEET_NAME)
    if dataframe is None or dataframe.empty:
        yield {"t": "log", "level": "ERROR",
               "msg": f"Sheet '{SHEET_NAME}' kosong atau tidak ada."}
        yield {"t": "done", "generated": [], "skipped": []}
        return

    dataframe = _populated_rows(dataframe.fillna(""))
    generated, skipped = [], []
    total = len(dataframe)
    yield {"t": "log", "level": "STEP",
           "msg": f"Memproses {total} dokumen SPTD..."}

    for done, (index, row) in enumerate(dataframe.iterrows(), start=1):
        excel_row = index + 2
        nama_ppl = clean_value(row.get("nama_ppl", ""))
        try:
            keluarga = _count_value(
                row.get("jml_keluarga", ""),
                row_number=excel_row,
                column="jml_keluarga",
            )
            usaha = _count_value(
                row.get("jml_usaha", ""),
                row_number=excel_row,
                column="jml_usaha",
            )
        except ValueError as exc:
            skipped.append(f"baris {excel_row} ({exc})")
            yield {"t": "log", "level": "WARN", "msg": f"   {exc} Dilewati."}
            yield {"t": "progress", "done": done, "total": total}
            continue

        replacements = row_placeholder_replacements(row)
        replacements["{{jml_keluarga}}"] = _format_decimal(keluarga)
        replacements["{{jml_usaha}}"] = _format_decimal(usaha)
        replacements["{{responden}}"] = _format_decimal(keluarga + usaha)
        document = Document(template_path)
        replace_text_preserving_runs(document, replacements)

        out_name = f"SPTD_{done:03d}_{_slug(nama_ppl)}.docx"
        out_path = os.path.join(out_dir, out_name)
        document.save(out_path)
        generated.append(out_path)
        yield {"t": "file", "path": out_path}
        yield {"t": "log", "level": "OK", "msg": f"   Tersimpan: {out_name}"}
        yield {"t": "progress", "done": done, "total": total}

    yield {"t": "log", "level": "OK" if generated else "ERROR",
           "msg": f"Selesai: {len(generated)} dokumen berhasil, "
                  f"{len(skipped)} dilewati."}
    yield {"t": "done", "generated": generated, "skipped": skipped}